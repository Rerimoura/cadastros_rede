import streamlit as st
import pandas as pd
import openpyxl
from copy import copy
from io import BytesIO
from datetime import datetime
from urllib.parse import quote_plus
from sqlalchemy import create_engine
import re

st.set_page_config(page_title="Cadastro de Produtos - Redes", page_icon="📋", layout="wide")


def col_mercadoria(p): return int_or_none(p['mercadoria'])
def col_ean(p): return int_or_none(p['extra'])
def col_descricao(p): return p['descricao']
def col_marca(p): return p['marca']
def col_funcao_un(p): return 'UN'
def col_cest(p): return int_or_none(p['cest'])
def col_ncm(p): return p['classificacao_fiscal']
def col_dun14(p): return int_or_none(p['dun14'])
def col_peso(p): return float_or_none(p['peso_unitario'])
def col_embalagem_frasco(p): return 'Frasco'


def col_qtd_embalagem(p):
    qtde = p['qtde']
    qtde_und = p['qtde_und']
    if qtde is None or qtde_und is None:
        return None
    return float(qtde) * float(qtde_und)


def col_validade(p):
    dias = p['dias_validade']
    if dias is None:
        return None
    # dias_validade vem em dias corridos; convertido para meses (base 30 dias/mês)
    meses = int(dias) // 30
    return f"{meses} meses"


# ── Configuração por rede ───────────────────────────────────────────────────
# Para adicionar uma nova rede: colocar o arquivo MODELO em templates/ e criar
# uma nova entrada aqui com o mapeamento de colunas dessa planilha.
REDE_CONFIGS = {
    'SuperMaxi': {
        'template_path': 'templates/supermaxi.xlsx',
        'sheet_name': 'Plan1',
        'linha_inicial': 9,
        'colunas': {
            'A': col_mercadoria,
            'B': col_ean,
            'C': col_descricao,
            'E': col_marca,
            'F': col_funcao_un,
            'G': col_cest,
            'H': col_ncm,
            'J': col_dun14,
            'M': col_peso,
            'N': col_peso,
            'O': col_embalagem_frasco,
            'P': col_qtd_embalagem,
            'Q': col_validade,
        },
    },
}


# ── Banco de dados ───────────────────────────────────────────────────────────

@st.cache_resource
def get_connection():
    """
    Cria um engine SQLAlchemy com pool_pre_ping=True.
    Isso testa a conexão antes de cada query e reconecta automaticamente
    se ela tiver caído (resolve o erro "connection already closed" causado
    por conexões ociosas derrubadas pelo servidor — o app fica minutos sem
    uso entre um cadastro e outro).
    """
    try:
        cfg = st.secrets["postgres"]
        senha = quote_plus(cfg["password"])  # evita erro se a senha tiver caracteres especiais
        url = (
            f"postgresql+psycopg2://{cfg['user']}:{senha}"
            f"@{cfg['host']}:{cfg['port']}/{cfg['database']}"
        )
        return create_engine(
            url,
            pool_pre_ping=True,   # testa a conexão antes de usar; reconecta se estiver morta
            pool_recycle=1800,    # recicla conexões a cada 30 min
        )
    except Exception as e:
        st.error(f"Erro ao conectar ao banco de dados: {e}")
        return None


@st.cache_data(ttl=600)
def buscar_produtos(_conn, codigos):
    query = """
        SELECT
            m.mercadoria, m.descricao, m.extra, m.dun14, m.classificacao_fiscal,
            m.cest, m.peso_unitario, m.qtde, m.qtde_und, m.dias_validade,
            c.marca
        FROM mercadorias m
        LEFT JOIN categorias c ON c.mercadoria = m.mercadoria
        WHERE m.mercadoria = ANY(%s)
    """
    try:
        return pd.read_sql(query, _conn, params=(codigos,))
    except Exception as e:
        st.error(f"Erro ao buscar produtos: {e}")
        return pd.DataFrame()


# ── Utilitários ───────────────────────────────────────────────────────────────

def int_or_none(val):
    if val is None or pd.isna(val):
        return None
    try:
        return int(val)
    except (ValueError, TypeError):
        return val


def float_or_none(val):
    if val is None or pd.isna(val):
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return val


def parse_codigos(texto):
    """Aceita códigos separados por quebra de linha, vírgula, ponto e vírgula
    ou espaço, preservando a ordem de digitação e sem repetir."""
    partes = re.split(r'[\s,;]+', texto.strip())
    codigos = []
    vistos = set()
    for parte in partes:
        parte = parte.strip()
        if not parte:
            continue
        try:
            codigo = int(parte)
        except ValueError:
            continue
        if codigo not in vistos:
            vistos.add(codigo)
            codigos.append(codigo)
    return codigos


def gerar_planilha(rede_config, produtos):
    """Preenche o template da rede a partir da linha_inicial, um produto por
    linha, copiando o estilo da linha-modelo para as linhas extras criadas."""
    wb = openpyxl.load_workbook(rede_config['template_path'])
    ws = wb[rede_config['sheet_name']]
    linha_modelo = rede_config['linha_inicial']

    for i, produto in enumerate(produtos):
        linha = linha_modelo + i
        if linha != linha_modelo:
            for col in range(1, ws.max_column + 1):
                origem = ws.cell(row=linha_modelo, column=col)
                destino = ws.cell(row=linha, column=col)
                destino.font = copy(origem.font)
                destino.border = copy(origem.border)
                destino.fill = copy(origem.fill)
                destino.number_format = origem.number_format
                destino.alignment = copy(origem.alignment)
            ws.row_dimensions[linha].height = ws.row_dimensions[linha_modelo].height

        for col_letra, funcao in rede_config['colunas'].items():
            ws[f'{col_letra}{linha}'] = funcao(produto)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ── App principal ─────────────────────────────────────────────────────────────

def main():
    st.title("📋 Cadastro de Produtos - Redes")
    st.caption(
        "Cole os códigos das mercadorias e receba a ficha de cadastro da rede "
        "já preenchida com os dados do banco."
    )
    st.markdown("---")

    conn = get_connection()
    if not conn:
        return

    rede = st.selectbox("Rede", list(REDE_CONFIGS.keys()))
    rede_config = REDE_CONFIGS[rede]

    texto_codigos = st.text_area(
        "Códigos das mercadorias",
        height=150,
        placeholder="Cole os códigos, um por linha ou separados por vírgula.\nEx: 306543, 208949, 308763"
    )

    if st.button("Gerar planilha", type="primary"):
        codigos = parse_codigos(texto_codigos)

        if not codigos:
            st.warning("Informe ao menos um código de mercadoria válido.")
            return

        with st.spinner("Buscando produtos no banco de dados..."):
            df = buscar_produtos(conn, codigos)

        encontrados = set(df['mercadoria'].astype(int)) if not df.empty else set()
        faltando = [c for c in codigos if c not in encontrados]

        if df.empty:
            st.error("Nenhum dos códigos informados foi encontrado na base.")
            return

        if faltando:
            st.warning(
                f"⚠️ {len(faltando)} código(s) não encontrado(s) na base e que "
                f"ficarão de fora da planilha: {', '.join(map(str, faltando))}"
            )

        # Mantém a ordem em que os códigos foram digitados/colados
        df['mercadoria'] = df['mercadoria'].astype(int)
        ordem = [c for c in codigos if c in encontrados]
        df = df.set_index('mercadoria').loc[ordem].reset_index()

        st.success(f"✅ {len(df)} produto(s) encontrado(s).")
        st.dataframe(df, use_container_width=True)

        buffer = gerar_planilha(rede_config, df.to_dict('records'))

        st.download_button(
            "📥 Baixar planilha preenchida",
            data=buffer,
            file_name=f"Cadastro_{rede}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary"
        )

    st.markdown("---")
    st.caption(
        "📌 Campos não listados no De/Para (Complemento, Forma de Aquisição, "
        "Tipo de Troca, Tributação) ficam em branco — são preenchidos manualmente."
    )


if __name__ == "__main__":
    main()
